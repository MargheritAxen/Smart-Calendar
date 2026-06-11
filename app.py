import streamlit as st
import pandas as pd
import holidays
import gspread
from google.oauth2.service_account import Credentials
from datetime import date, timedelta
from io import BytesIO
import calendar
import streamlit.components.v1 as components

st.set_page_config(
    page_title="Smart Calendar",
    page_icon="📅",
    layout="centered"
)

# ===== CONFIG =====
STATI = ["Ufficio", "Smart", "Ferie", "Malattia", "Jol"]
PERSONE = ["Margherita", "Roberto"]

NOME_FILE_GOOGLE_SHEETS = "Smart Calendar"

FOGLIO_PRESENZE = "presenze"
FOGLIO_FESTE = "festivita_manuali"
FOGLIO_RIEPILOGO = "riepilogo"

festivita_italiane = holidays.Italy()

CODICI = {
    "Ufficio": "UFF",
    "Smart": "LAW",
    "Ferie": "FER",
    "Malattia": "MAL",
    "Jol": "JOL",
    "Assenza": "FER",
}

COLORI = {
    "UFF": "#ff8a00",
    "LAW": "#19e635",
    "FER": "#8a078a",
    "MAL": "#fff200",
    "JOL": "#6f42c1",
}

st.markdown("""
<style>
.block-container {
    padding-top: 5.5rem;
    max-width: 820px;
}

.smart-title {
    font-size: 2rem;
    font-weight: 800;
    margin-bottom: 0.8rem;
    line-height: 1.2;
}

.legend-compact {
    display: grid;
    grid-template-columns: repeat(5, 1fr);
    gap: 4px;
    margin: 6px 0 8px 0;
}

.legend-item {
    border-radius: 8px;
    overflow: hidden;
    border: 1px solid rgba(255,255,255,0.18);
    text-align: center;
    font-size: 0.62rem;
}

.legend-code {
    font-weight: 900;
    padding: 4px 2px;
    color: #000;
}

.legend-desc {
    padding: 4px 2px;
    background: rgba(255,255,255,0.05);
    font-size: 0.58rem;
    min-height: 20px;
}

.month-title {
    font-size: 1.45rem;
    font-weight: 900;
    margin: 22px 0 10px 0;
}

.calendar-table {
    width: 100%;
    border-collapse: collapse;
    table-layout: fixed;
    font-size: 0.78rem;
}

.calendar-table th {
    padding: 8px 4px;
    background: rgba(255,255,255,0.08);
    text-align: left;
    font-weight: 800;
    color: rgba(255,255,255,0.75);
}

.calendar-table td {
    height: 78px;
    vertical-align: top;
    padding: 6px;
    border: 1px solid rgba(255,255,255,0.12);
    font-weight: 800;
}

.day-number {
    font-size: 0.78rem;
    margin-bottom: 4px;
}

.pres-row {
    display: block;
    font-size: 0.62rem;
    line-height: 1.1rem;
    white-space: nowrap;
}

@media (max-width: 600px) {
    .legend-compact {
        grid-template-columns: repeat(3, 1fr);
    }

    .smart-title {
        font-size: 1.7rem;
    }

    .calendar-table {
        font-size: 0.68rem;
    }

    .calendar-table td {
        height: 70px;
        padding: 4px;
    }

    .pres-row {
        font-size: 0.62rem;
        line-height: 0.95rem;
    }
}
</style>
""", unsafe_allow_html=True)


@st.cache_resource
def connetti_google_sheets():
    scope = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive"
    ]

    credentials = Credentials.from_service_account_info(
        st.secrets["gcp_service_account"],
        scopes=scope
    )

    client = gspread.authorize(credentials)

    try:
        spreadsheet = client.open(NOME_FILE_GOOGLE_SHEETS)
    except gspread.SpreadsheetNotFound:
        spreadsheet = client.create(NOME_FILE_GOOGLE_SHEETS)

    return spreadsheet


@st.cache_resource
def inizializza_fogli():
    spreadsheet = connetti_google_sheets()

    fogli_esistenti = {ws.title: ws for ws in spreadsheet.worksheets()}

    def crea_o_prendi(nome_foglio, intestazioni):
        if nome_foglio in fogli_esistenti:
            worksheet = fogli_esistenti[nome_foglio]
        else:
            worksheet = spreadsheet.add_worksheet(
                title=nome_foglio,
                rows=1000,
                cols=len(intestazioni)
            )
            worksheet.update(values=[intestazioni], range_name="A1")

        valori = worksheet.get_all_values()
        if len(valori) == 0:
            worksheet.update(values=[intestazioni], range_name="A1")

        return worksheet

    ws_presenze = crea_o_prendi(FOGLIO_PRESENZE, ["data", "persona", "giustificativo"])
    ws_feste = crea_o_prendi(FOGLIO_FESTE, ["data", "descrizione"])
    ws_riepilogo = crea_o_prendi(
        FOGLIO_RIEPILOGO,
        [
            "persona", "trimestre", "UFF", "LAW", "FER", "MAL", "JOL",
            "Giorni conteggiati", "% UFF", "% LAW", "Esito"
        ]
    )

    try:
        fogli_aggiornati = spreadsheet.worksheets()
        if len(fogli_aggiornati) > 1:
            for ws in fogli_aggiornati:
                if ws.title == "Foglio1":
                    spreadsheet.del_worksheet(ws)
                    break
    except Exception:
        pass

    return spreadsheet, ws_presenze, ws_feste, ws_riepilogo


spreadsheet, ws_presenze, ws_feste, ws_riepilogo = inizializza_fogli()


def feste_extra_aziendali(anno):
    from dateutil.easter import easter

    pasqua = easter(anno)

    return {
        date(anno, 3, 19),
        pasqua + timedelta(days=39),
        pasqua + timedelta(days=60),
        date(anno, 11, 4),
        date(anno, 12, 8),
        date(anno, 11, 2),
        date(anno, 12, 24),
        date(anno, 12, 31),
    }


def date_festive_manuali(df_feste):
    if df_feste.empty:
        return set()

    tmp = df_feste.copy()
    tmp["data"] = pd.to_datetime(tmp["data"], errors="coerce")
    tmp = tmp.dropna(subset=["data"])
    return set(tmp["data"].dt.date)


def is_giorno_bloccato(giorno, df_feste):
    giorno = pd.to_datetime(giorno).date()
    feste_extra = feste_extra_aziendali(giorno.year)
    feste_manuali = date_festive_manuali(df_feste)

    return (
        giorno.weekday() >= 5
        or giorno in festivita_italiane
        or giorno in feste_extra
        or giorno in feste_manuali
    )


def codice_giustificativo(giustificativo):
    return CODICI.get(giustificativo, giustificativo)


@st.cache_data(ttl=30)
def leggi_presenze_cache():
    dati = ws_presenze.get_all_records()

    if not dati:
        return pd.DataFrame(columns=["data", "persona", "giustificativo"])

    df = pd.DataFrame(dati)
    df["data"] = pd.to_datetime(df["data"], errors="coerce")
    df = df.dropna(subset=["data"])
    return df


@st.cache_data(ttl=30)
def leggi_feste_cache():
    dati = ws_feste.get_all_records()

    if not dati:
        return pd.DataFrame(columns=["data", "descrizione"])

    df = pd.DataFrame(dati)
    df["data"] = pd.to_datetime(df["data"], errors="coerce")
    df = df.dropna(subset=["data"])
    return df


def leggi_presenze():
    return leggi_presenze_cache()


def leggi_feste_manuali():
    return leggi_feste_cache()


def svuota_cache_dati():
    leggi_presenze_cache.clear()
    leggi_feste_cache.clear()


def salva_presenze(df):
    df = df.copy()
    df["data"] = pd.to_datetime(df["data"]).dt.strftime("%Y-%m-%d")

    ws_presenze.clear()
    ws_presenze.update(values=[["data", "persona", "giustificativo"]], range_name="A1")

    if not df.empty:
        ws_presenze.append_rows(df.values.tolist())

    svuota_cache_dati()


def salva_feste_manuali(df):
    df = df.copy()
    df["data"] = pd.to_datetime(df["data"]).dt.strftime("%Y-%m-%d")

    ws_feste.clear()
    ws_feste.update(values=[["data", "descrizione"]], range_name="A1")

    if not df.empty:
        ws_feste.append_rows(df.values.tolist())

    svuota_cache_dati()


def salva_riepilogo(riepilogo):
    intestazioni = [
        "persona", "trimestre", "UFF", "LAW", "FER", "MAL", "JOL",
        "Giorni conteggiati", "% UFF", "% LAW", "Esito"
    ]

    ws_riepilogo.clear()
    ws_riepilogo.update(values=[intestazioni], range_name="A1")

    if not riepilogo.empty:
        ws_riepilogo.append_rows(riepilogo[intestazioni].values.tolist())


IMPORT_PRESENZE_MARGHERITA = [('2026-01-02', 'Margherita', 'Smart'), ('2026-01-05', 'Margherita', 'Ufficio'), ('2026-01-07', 'Margherita', 'Ufficio'), ('2026-01-08', 'Margherita', 'Smart'), ('2026-01-09', 'Margherita', 'Ufficio'), ('2026-01-12', 'Margherita', 'Smart'), ('2026-01-13', 'Margherita', 'Ufficio'), ('2026-01-14', 'Margherita', 'Ufficio'), ('2026-01-15', 'Margherita', 'Assenza'), ('2026-01-16', 'Margherita', 'Smart'), ('2026-01-19', 'Margherita', 'Assenza'), ('2026-01-20', 'Margherita', 'Smart'), ('2026-01-21', 'Margherita', 'Ufficio'), ('2026-01-22', 'Margherita', 'Ufficio'), ('2026-01-23', 'Margherita', 'Ufficio'), ('2026-01-26', 'Margherita', 'Ufficio'), ('2026-01-27', 'Margherita', 'Smart'), ('2026-01-28', 'Margherita', 'Ufficio'), ('2026-01-29', 'Margherita', 'Smart'), ('2026-01-30', 'Margherita', 'Smart'), ('2026-02-02', 'Margherita', 'Smart'), ('2026-02-03', 'Margherita', 'Assenza'), ('2026-02-04', 'Margherita', 'Smart'), ('2026-02-05', 'Margherita', 'Smart'), ('2026-02-06', 'Margherita', 'Assenza'), ('2026-02-09', 'Margherita', 'Smart'), ('2026-02-10', 'Margherita', 'Ufficio'), ('2026-02-11', 'Margherita', 'Smart'), ('2026-02-12', 'Margherita', 'Ufficio'), ('2026-02-13', 'Margherita', 'Ufficio'), ('2026-02-16', 'Margherita', 'Ufficio'), ('2026-02-17', 'Margherita', 'Smart'), ('2026-02-18', 'Margherita', 'Smart'), ('2026-02-19', 'Margherita', 'Ufficio'), ('2026-02-20', 'Margherita', 'Smart'), ('2026-02-23', 'Margherita', 'Ufficio'), ('2026-02-24', 'Margherita', 'Smart'), ('2026-02-25', 'Margherita', 'Ufficio'), ('2026-02-26', 'Margherita', 'Ufficio'), ('2026-02-27', 'Margherita', 'Smart'), ('2026-03-02', 'Margherita', 'Ufficio'), ('2026-03-03', 'Margherita', 'Smart'), ('2026-03-04', 'Margherita', 'Smart'), ('2026-03-05', 'Margherita', 'Smart'), ('2026-03-06', 'Margherita', 'Ufficio'), ('2026-03-09', 'Margherita', 'Ufficio'), ('2026-03-10', 'Margherita', 'Smart'), ('2026-03-11', 'Margherita', 'Ufficio'), ('2026-03-12', 'Margherita', 'Ufficio'), ('2026-03-13', 'Margherita', 'Ufficio'), ('2026-03-16', 'Margherita', 'Ufficio'), ('2026-03-17', 'Margherita', 'Smart'), ('2026-03-18', 'Margherita', 'Smart'), ('2026-03-19', 'Margherita', 'Assenza'), ('2026-03-20', 'Margherita', 'Assenza'), ('2026-03-23', 'Margherita', 'Assenza'), ('2026-03-24', 'Margherita', 'Assenza'), ('2026-03-25', 'Margherita', 'Assenza'), ('2026-03-26', 'Margherita', 'Assenza'), ('2026-03-27', 'Margherita', 'Smart'), ('2026-03-30', 'Margherita', 'Smart'), ('2026-03-31', 'Margherita', 'Smart'), ('2026-04-01', 'Margherita', 'Smart'), ('2026-04-02', 'Margherita', 'Assenza'), ('2026-04-07', 'Margherita', 'Assenza'), ('2026-04-08', 'Margherita', 'Assenza'), ('2026-04-09', 'Margherita', 'Ufficio'), ('2026-04-10', 'Margherita', 'Ufficio'), ('2026-04-13', 'Margherita', 'Ufficio'), ('2026-04-14', 'Margherita', 'Ufficio'), ('2026-04-15', 'Margherita', 'Smart'), ('2026-04-16', 'Margherita', 'Ufficio'), ('2026-04-17', 'Margherita', 'Smart'), ('2026-04-20', 'Margherita', 'Ufficio'), ('2026-04-21', 'Margherita', 'Smart'), ('2026-04-22', 'Margherita', 'Ufficio'), ('2026-04-23', 'Margherita', 'Smart'), ('2026-04-24', 'Margherita', 'Smart'), ('2026-04-27', 'Margherita', 'Ufficio'), ('2026-04-28', 'Margherita', 'Smart'), ('2026-04-29', 'Margherita', 'Ufficio'), ('2026-04-30', 'Margherita', 'Smart'), ('2026-05-04', 'Margherita', 'Ufficio'), ('2026-05-05', 'Margherita', 'Smart'), ('2026-05-06', 'Margherita', 'Smart'), ('2026-05-07', 'Margherita', 'Smart'), ('2026-05-08', 'Margherita', 'Ufficio'), ('2026-05-11', 'Margherita', 'Ufficio'), ('2026-05-12', 'Margherita', 'Smart'), ('2026-05-13', 'Margherita', 'Ufficio'), ('2026-05-14', 'Margherita', 'Smart'), ('2026-05-15', 'Margherita', 'Ufficio'), ('2026-05-18', 'Margherita', 'Assenza'), ('2026-05-19', 'Margherita', 'Assenza'), ('2026-05-20', 'Margherita', 'Assenza'), ('2026-05-21', 'Margherita', 'Assenza'), ('2026-05-22', 'Margherita', 'Smart'), ('2026-05-25', 'Margherita', 'Ufficio'), ('2026-05-26', 'Margherita', 'Assenza'), ('2026-05-27', 'Margherita', 'Smart'), ('2026-05-28', 'Margherita', 'Smart'), ('2026-05-29', 'Margherita', 'Assenza'), ('2026-06-01', 'Margherita', 'Assenza'), ('2026-06-03', 'Margherita', 'Ufficio'), ('2026-06-04', 'Margherita', 'Smart'), ('2026-06-05', 'Margherita', 'Smart'), ('2026-06-08', 'Margherita', 'Smart'), ('2026-06-09', 'Margherita', 'Smart')]


def importa_presenze_margherita_una_volta():
    valori = ws_presenze.get_all_records()
    esistenti = set()

    for r in valori:
        data_raw = r.get("data")
        persona = str(r.get("persona", "")).strip()
        if data_raw and persona:
            data_norm = pd.to_datetime(data_raw, errors="coerce")
            if pd.notna(data_norm):
                esistenti.add((data_norm.strftime("%Y-%m-%d"), persona))

    righe_da_aggiungere = []
    for data_str, persona, giustificativo in IMPORT_PRESENZE_MARGHERITA:
        chiave = (data_str, persona)
        if chiave not in esistenti:
            righe_da_aggiungere.append([data_str, persona, giustificativo])

    if righe_da_aggiungere:
        ws_presenze.append_rows(righe_da_aggiungere)
        svuota_cache_dati()

    return len(righe_da_aggiungere)


def calcola_riepilogo(df, df_feste):
    if df.empty:
        return pd.DataFrame(), pd.DataFrame()

    df = df.copy()
    df["data"] = pd.to_datetime(df["data"], errors="coerce")
    df = df.dropna(subset=["data"])

    date_feste_manuali = date_festive_manuali(df_feste)

    date_feste_extra = set()
    for anno in df["data"].dt.year.dropna().unique():
        date_feste_extra.update(feste_extra_aziendali(int(anno)))

    df["weekend"] = df["data"].dt.weekday >= 5
    df["festivo_italia"] = df["data"].dt.date.apply(lambda x: x in festivita_italiane)
    df["festivo_manuale"] = df["data"].dt.date.apply(lambda x: x in date_feste_manuali)
    df["festivo_extra"] = df["data"].dt.date.apply(lambda x: x in date_feste_extra)

    df_valido = df[
        (~df["weekend"]) &
        (~df["festivo_italia"]) &
        (~df["festivo_manuale"]) &
        (~df["festivo_extra"])
    ].copy()

    if df_valido.empty:
        return pd.DataFrame(), df_valido

    df_valido["trimestre"] = df_valido["data"].dt.to_period("Q").astype(str)
    df_valido["codice"] = df_valido["giustificativo"].map(CODICI).fillna(df_valido["giustificativo"])

    riepilogo = (
        df_valido
        .groupby(["persona", "trimestre", "codice"])
        .size()
        .unstack(fill_value=0)
        .reset_index()
    )

    for col in ["UFF", "LAW", "FER", "MAL", "JOL"]:
        if col not in riepilogo.columns:
            riepilogo[col] = 0

    riepilogo["Giorni conteggiati"] = riepilogo["UFF"] + riepilogo["LAW"]

    riepilogo["% UFF"] = (
        riepilogo["UFF"] / riepilogo["Giorni conteggiati"] * 100
    ).replace([float("inf"), -float("inf")], 0).fillna(0).round(2)

    riepilogo["% LAW"] = (
        riepilogo["LAW"] / riepilogo["Giorni conteggiati"] * 100
    ).replace([float("inf"), -float("inf")], 0).fillna(0).round(2)

    riepilogo["Esito"] = riepilogo.apply(
        lambda r: "OK" if r["Giorni conteggiati"] > 0 and 40 <= r["% UFF"] <= 60 else "KO",
        axis=1
    )

    ordine_colonne = [
        "persona", "trimestre", "UFF", "LAW", "FER", "MAL", "JOL",
        "Giorni conteggiati", "% UFF", "% LAW", "Esito"
    ]

    return riepilogo[ordine_colonne], df_valido


def render_legenda():
    st.markdown("""
    <div class="legend-compact">
        <div class="legend-item">
            <div class="legend-code" style="background:#ff8a00;">UFF</div>
            <div class="legend-desc">Ufficio</div>
        </div>
        <div class="legend-item">
            <div class="legend-code" style="background:#19e635;">LAW</div>
            <div class="legend-desc">Smart</div>
        </div>
        <div class="legend-item">
            <div class="legend-code" style="background:#8a078a;color:white;">FER</div>
            <div class="legend-desc">Ferie</div>
        </div>
        <div class="legend-item">
            <div class="legend-code" style="background:#fff200;">MAL</div>
            <div class="legend-desc">Malattia</div>
        </div>
        <div class="legend-item">
            <div class="legend-code" style="background:#6f42c1;color:white;">JOL</div>
            <div class="legend-desc">Jolly</div>
        </div>
    </div>
    """, unsafe_allow_html=True)
def abbrevia_persona(nome):
    nome = str(nome).strip()
    if nome.lower().startswith("margherita"):
        return "MAR"
    if nome.lower().startswith("roberto"):
        return "ROB"
    return nome[:3].upper()


def render_calendario_mese(df, df_feste, anno, mese):
    'Calendario del mese corrente: mostra solo MAR/ROB colorati in base al giustificativo.'
    mesi_it = [
        'GENNAIO', 'FEBBRAIO', 'MARZO', 'APRILE', 'MAGGIO', 'GIUGNO',
        'LUGLIO', 'AGOSTO', 'SETTEMBRE', 'OTTOBRE', 'NOVEMBRE', 'DICEMBRE'
    ]

    giorni = ['Lun', 'Mar', 'Mer', 'Gio', 'Ven', 'Sab', 'Dom']

    lookup = {}
    if not df.empty:
        tmp = df.copy()
        tmp['data'] = pd.to_datetime(tmp['data'], errors='coerce').dt.date
        tmp = tmp.dropna(subset=['data'])
        tmp['codice'] = tmp['giustificativo'].map(CODICI).fillna(tmp['giustificativo'])

        for _, r in tmp.iterrows():
            lookup[(r['data'], r['persona'])] = r['codice']

    colori = {
        'UFF': '#ff8a00',
        'LAW': '#19e635',
        'FER': '#8a078a',
        'MAL': '#fff200',
        'JOL': '#6f42c1',
        'EMPTY': '#111827',
    }

    testo_colore = {
        'UFF': '#000000',
        'LAW': '#000000',
        'FER': '#ffffff',
        'MAL': '#000000',
        'JOL': '#ffffff',
    }

    _, giorni_mese = calendar.monthrange(anno, mese)

    settimane = []
    settimana = [None for _ in range(7)]

    for day in range(1, giorni_mese + 1):
        data_giorno = date(anno, mese, day)
        weekday = data_giorno.weekday()
        bloccato = is_giorno_bloccato(data_giorno, df_feste)

        if bloccato:
            contenuto = f'<div class="day-number">{day}</div>'
            colore_bg = colori['EMPTY']
            colore_txt = '#ffffff'
        else:
            righe = [f'<div class="day-number">{day}</div>']

            for persona in PERSONE:
                codice = lookup.get((data_giorno, persona), '')
                if codice:
                    nome_breve = abbrevia_persona(persona)
                    colore_nome = colori.get(codice, colori['EMPTY'])
                    colore_testo = testo_colore.get(codice, '#ffffff')
                    righe.append(
                        f'<div class="name-chip" style="background:{colore_nome}; color:{colore_testo};">{nome_breve}</div>'
                    )

            contenuto = ''.join(righe)
            colore_bg = colori['EMPTY']
            colore_txt = '#ffffff'

        cella = f'<td style="background:{colore_bg}; color:{colore_txt};">{contenuto}</td>'
        settimana[weekday] = cella

        if weekday == 6:
            settimane.append(settimana)
            settimana = [None for _ in range(7)]

    if any(x is not None for x in settimana):
        settimane.append(settimana)

    intestazione_giorni = ''.join([f'<th>{g}</th>' for g in giorni])

    html = f'''
    <style>
        body {{
            margin: 0;
            background: transparent;
            font-family: Arial, sans-serif;
            color: white;
        }}

        .month-title {{
            font-size: 26px;
            font-weight: 900;
            margin: 0 0 12px 0;
            color: white;
        }}

        .calendar-table {{
            width: 100%;
            border-collapse: collapse;
            table-layout: fixed;
            overflow: hidden;
            border-radius: 8px;
        }}

        .calendar-table th {{
            padding: 9px 4px;
            background: #1f2937;
            color: #d1d5db;
            text-align: center;
            font-size: 13px;
            border: 1px solid #374151;
        }}

        .calendar-table td {{
            height: 82px;
            vertical-align: top;
            padding: 6px;
            border: 1px solid #374151;
            box-sizing: border-box;
        }}

        .day-number {{
            font-size: 13px;
            font-weight: 900;
            margin-bottom: 7px;
        }}

        .dark-number {{
            color: #000000;
        }}

        .name-chip {{
            display: block;
            width: 100%;
            border-radius: 8px;
            text-align: center;
            font-size: 12px;
            font-weight: 900;
            line-height: 18px;
            margin-bottom: 5px;
            border: 1px solid rgba(0,0,0,0.28);
            box-sizing: border-box;
        }}

        @media (max-width: 600px) {{
            .month-title {{
                font-size: 22px;
            }}

            .calendar-table th {{
                font-size: 11px;
                padding: 7px 2px;
            }}

            .calendar-table td {{
                height: 72px;
                padding: 4px;
            }}

            .day-number {{
                font-size: 11px;
                margin-bottom: 5px;
            }}

            .name-chip {{
                font-size: 10px;
                line-height: 15px;
                margin-bottom: 4px;
                border-radius: 6px;
            }}
        }}
    </style>

    <div class="month-title">{mesi_it[mese - 1]} {anno}</div>

    <table class="calendar-table">
        <thead>
            <tr>{intestazione_giorni}</tr>
        </thead>
        <tbody>
    '''

    for settimana in settimane:
        html += '<tr>'
        for cella in settimana:
            html += cella if cella is not None else '<td></td>'
        html += '</tr>'

    html += '''
        </tbody>
    </table>
    '''

    components.html(html, height=560, scrolling=False)

def genera_excel_formattato(df, df_feste, riepilogo, anno):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    buffer = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = str(anno)
    ws_riep = wb.create_sheet("Riepilogo")

    persone = PERSONE
    mesi = [
        "GENNAIO", "FEBBRAIO", "MARZO", "APRILE", "MAGGIO", "GIUGNO",
        "LUGLIO", "AGOSTO", "SETTEMBRE", "OTTOBRE", "NOVEMBRE", "DICEMBRE"
    ]
    giorni_it = ["Lu", "Ma", "Me", "Gi", "Ve", "Sa", "Do"]

    feste_extra = feste_extra_aziendali(anno)
    feste_manuali = date_festive_manuali(df_feste)

    fill_header = PatternFill("solid", fgColor="D9EAF7")
    fill_month = PatternFill("solid", fgColor="1F4E78")
    fill_uff = PatternFill("solid", fgColor="F9CB9C")
    fill_law = PatternFill("solid", fgColor="B6D7A8")
    fill_fer = PatternFill("solid", fgColor="D5A6BD")
    fill_mal = PatternFill("solid", fgColor="FFF2CC")
    fill_jol = PatternFill("solid", fgColor="D9D2E9")

    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    presenze = df.copy()
    if not presenze.empty:
        presenze["data"] = pd.to_datetime(presenze["data"]).dt.date
        presenze["codice"] = presenze["giustificativo"].map(CODICI).fillna(presenze["giustificativo"])
        lookup = {(r["data"], r["persona"]): r["codice"] for _, r in presenze.iterrows()}
    else:
        lookup = {}

    ws["A1"] = "LEGENDA"
    ws["A1"].font = Font(bold=True)
    legenda = [
        ("UFF", "Presenza"),
        ("LAW", "Smart working"),
        ("FER", "Ferie - non conta"),
        ("MAL", "Malattia - non conta"),
        ("JOL", "Jolly - non conta"),
    ]
    for i, (codice, descrizione) in enumerate(legenda, start=2):
        ws[f"A{i}"] = codice
        ws[f"B{i}"] = descrizione
        ws[f"A{i}"].font = Font(bold=True)
        ws[f"A{i}"].border = border
        ws[f"B{i}"].border = border

    start_cols = [1, 6, 11]
    start_rows = [10, 47, 84, 121]

    codice_fill = {
        "UFF": fill_uff,
        "LAW": fill_law,
        "FER": fill_fer,
        "MAL": fill_mal,
        "JOL": fill_jol,
    }

    for month in range(1, 13):
        block_col = start_cols[(month - 1) % 3]
        block_row = start_rows[(month - 1) // 3]

        ws.merge_cells(start_row=block_row, start_column=block_col, end_row=block_row, end_column=block_col + 3)
        c = ws.cell(block_row, block_col, mesi[month - 1])
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill_month
        c.alignment = Alignment(horizontal="center")

        headers = ["Giorno", "Sett."] + persone
        for j, h in enumerate(headers):
            cell = ws.cell(block_row + 1, block_col + j, h)
            cell.font = Font(bold=True)
            cell.fill = fill_header
            cell.border = border
            cell.alignment = Alignment(horizontal="center")

        days = calendar.monthrange(anno, month)[1]
        for day in range(1, days + 1):
            row = block_row + 1 + day
            data = date(anno, month, day)
            weekday = data.weekday()
            is_festivo = data in festivita_italiane or data in feste_extra or data in feste_manuali

            values = [day, giorni_it[weekday]]
            for persona in persone:
                values.append("" if weekday >= 5 or is_festivo else lookup.get((data, persona), ""))

            for j, value in enumerate(values):
                cell = ws.cell(row, block_col + j, value)
                cell.border = border
                cell.alignment = Alignment(horizontal="center")
                if value in codice_fill:
                    cell.fill = codice_fill[value]
                    cell.font = Font(bold=True, color="000000")

    dv = DataValidation(type="list", formula1='"UFF,LAW,FER,MAL,JOL"', allow_blank=True)
    ws.add_data_validation(dv)

    for month in range(1, 13):
        block_col = start_cols[(month - 1) % 3]
        block_row = start_rows[(month - 1) // 3]
        days = calendar.monthrange(anno, month)[1]

        for day in range(1, days + 1):
            data = date(anno, month, day)
            if is_giorno_bloccato(data, df_feste):
                continue

            excel_row = block_row + 1 + day
            for p_idx in range(len(persone)):
                cell_ref = f"{get_column_letter(block_col + 2 + p_idx)}{excel_row}"
                dv.add(cell_ref)

    for col in range(1, 16):
        ws.column_dimensions[get_column_letter(col)].width = 13

    ws_riep["A1"] = "RIEPILOGO PRESENZE"
    ws_riep["A1"].font = Font(bold=True, size=14)

    headers = ["persona", "trimestre", "UFF", "LAW", "FER", "MAL", "JOL", "Giorni conteggiati", "% UFF", "% LAW", "Esito"]
    for j, h in enumerate(headers, start=1):
        cell = ws_riep.cell(3, j, h)
        cell.font = Font(bold=True)
        cell.fill = fill_header
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

    if not riepilogo.empty:
        for i, (_, r) in enumerate(riepilogo.iterrows(), start=4):
            for j, h in enumerate(headers, start=1):
                value = r[h] if h in r.index else ""
                cell = ws_riep.cell(i, j, value)
                cell.border = border
                cell.alignment = Alignment(horizontal="center")

    for col in range(1, 12):
        ws_riep.column_dimensions[get_column_letter(col)].width = 16

    wb.save(buffer)
    return buffer.getvalue()


def genera_pdf_presenze(df, df_feste, anno):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=18,
        leftMargin=18,
        topMargin=18,
        bottomMargin=18,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("MonthTitle", parent=styles["Title"], alignment=TA_CENTER, fontSize=26, leading=30)
    small_style = ParagraphStyle("Small", parent=styles["Normal"], fontSize=7, leading=8, alignment=TA_CENTER)

    elementi = []

    mesi = [
        "GENNAIO", "FEBBRAIO", "MARZO", "APRILE", "MAGGIO", "GIUGNO",
        "LUGLIO", "AGOSTO", "SETTEMBRE", "OTTOBRE", "NOVEMBRE", "DICEMBRE"
    ]
    giorni = ["Lun", "Mar", "Mer", "Gio", "Ven", "Sab", "Dom"]

    colore = {
        "UFF": colors.HexColor("#F9CB9C"),
        "LAW": colors.HexColor("#B6D7A8"),
        "FER": colors.HexColor("#D5A6BD"),
        "MAL": colors.HexColor("#FFF2CC"),
        "JOL": colors.HexColor("#D9D2E9"),
        "HEADER": colors.HexColor("#EEEEEE"),
    }

    presenze = df.copy()
    if not presenze.empty:
        presenze["data"] = pd.to_datetime(presenze["data"]).dt.date
        presenze["codice"] = presenze["giustificativo"].map(CODICI).fillna(presenze["giustificativo"])
        lookup = {(r["data"], r["persona"]): r["codice"] for _, r in presenze.iterrows()}
    else:
        lookup = {}

    for month in range(1, 13):
        elementi.append(Paragraph(f"{mesi[month - 1]} {anno}", title_style))
        elementi.append(Spacer(1, 6))

        legenda_pdf = "   ".join([
            "UFF = Presenza", "LAW = Smart", "FER = Ferie",
            "MAL = Malattia", "JOL = Jolly"
        ])
        elementi.append(Paragraph(legenda_pdf, small_style))
        elementi.append(Spacer(1, 8))

        dati = [giorni]
        celle_info = []

        _, days_in_month = calendar.monthrange(anno, month)
        row = ["" for _ in range(7)]
        row_info = [None for _ in range(7)]

        for day in range(1, days_in_month + 1):
            data = date(anno, month, day)
            wd = data.weekday()

            if is_giorno_bloccato(data, df_feste):
                testo = f"<b>{day}</b>"
                code_for_bg = ""
            else:
                righe = [f"<b>{day}</b>"]
                codici_presenti = []
                for persona in PERSONE:
                    codice = lookup.get((data, persona), "")
                    codici_presenti.append(codice)
                    if codice:
                        nome_breve = persona[:3]
                        righe.append(f"{nome_breve}: {codice}")
                testo = "<br/>".join(righe)
                code_for_bg = codici_presenti[0] if codici_presenti and codici_presenti[0] else ""

            row[wd] = Paragraph(testo, small_style)
            row_info[wd] = code_for_bg

            if wd == 6 or day == days_in_month:
                dati.append(row)
                celle_info.append(row_info)
                row = ["" for _ in range(7)]
                row_info = [None for _ in range(7)]

        tabella = Table(dati, colWidths=[108] * 7, rowHeights=[24] + [72] * (len(dati) - 1))
        stile = TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colore["HEADER"]),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
        ])

        for r_idx, info_row in enumerate(celle_info, start=1):
            for c_idx, codice in enumerate(info_row):
                if codice in colore:
                    stile.add("BACKGROUND", (c_idx, r_idx), (c_idx, r_idx), colore[codice])

        tabella.setStyle(stile)
        elementi.append(tabella)

        if month < 12:
            elementi.append(PageBreak())

    doc.build(elementi)
    return buffer.getvalue()


righe_importate_margherita = importa_presenze_margherita_una_volta()

df = leggi_presenze()
df_feste = leggi_feste_manuali()
riepilogo, df_valido = calcola_riepilogo(df, df_feste)


st.markdown('<div class="smart-title">📅 Smart Calendar</div>', unsafe_allow_html=True)

if righe_importate_margherita > 0:
    st.success(f"Import automatico completato: {righe_importate_margherita} presenze storiche di Margherita caricate.")

tab1, tab2, tab3 = st.tabs(["Inserisci", "Festività", "Riepilogo"])


with tab1:
    st.subheader("Nuova giornata")

    giorno = st.date_input("Data", value=date.today())
    persona = st.selectbox("Persona", PERSONE)
    giustificativo = st.selectbox("giustificativo", STATI, format_func=lambda x: f"{codice_giustificativo(x)} — {x}")

    if is_giorno_bloccato(giorno, df_feste):
        st.warning("Questo giorno è weekend o festività. Non va compilato e non entra nel conteggio.")
    else:
        if st.button("Salva presenza", width="stretch"):
            nuova_riga = pd.DataFrame([{"data": pd.to_datetime(giorno), "persona": persona, "giustificativo": giustificativo}])

            df = df[~((df["data"] == pd.to_datetime(giorno)) & (df["persona"] == persona))]
            df = pd.concat([df, nuova_riga], ignore_index=True)
            df = df.sort_values(["data", "persona"])

            salva_presenze(df)

            riepilogo, df_valido = calcola_riepilogo(df, df_feste)
            salva_riepilogo(riepilogo)

            st.success("Presenza salvata")
            st.rerun()

    if st.button("Elimina presenza selezionata", width="stretch"):
        df = df[~((df["data"] == pd.to_datetime(giorno)) & (df["persona"] == persona))]

        salva_presenze(df)

        riepilogo, df_valido = calcola_riepilogo(df, df_feste)
        salva_riepilogo(riepilogo)

        st.success("Presenza eliminata")
        st.rerun()

    st.divider()

    # Calendario consultabile mese per mese
    oggi = date.today()

    if "cal_anno" not in st.session_state:
        st.session_state.cal_anno = oggi.year

    if "cal_mese" not in st.session_state:
        st.session_state.cal_mese = oggi.month

    def cambia_mese(delta):
        nuovo_mese = st.session_state.cal_mese + delta
        nuovo_anno = st.session_state.cal_anno

        if nuovo_mese < 1:
            nuovo_mese = 12
            nuovo_anno -= 1

        if nuovo_mese > 12:
            nuovo_mese = 1
            nuovo_anno += 1

        st.session_state.cal_mese = nuovo_mese
        st.session_state.cal_anno = nuovo_anno

    st.markdown(
    """
    <div style="
        font-size:0.90rem;
        font-weight:700;
        color:#d1d5db;
        margin-bottom:6px;
        margin-top:4px;
    ">
        📖 LEGENDA
    </div>
    """,
    unsafe_allow_html=True
    )

    render_legenda()

    nav_action = st.query_params.get("nav")

    if nav_action == "prev":
        cambia_mese(-1)
        st.query_params.clear()
        st.rerun()

    if nav_action == "today":
        st.session_state.cal_anno = oggi.year
        st.session_state.cal_mese = oggi.month
        st.query_params.clear()
        st.rerun()

    if nav_action == "next":
        cambia_mese(1)
        st.query_params.clear()
        st.rerun()

    st.markdown("<div style='height:10px'></div>", unsafe_allow_html=True)

    st.markdown(
        """
        <style>
        .month-nav-row {
            display: flex;
            flex-direction: row;
            gap: 6px;
            width: 100%;
            margin: 8px 0 12px 0;
        }

        .month-nav-row a {
            flex: 1;
            display: block;
            text-align: center;
            text-decoration: none;
            color: #ffffff !important;
            background: #111827;
            border: 1px solid #4b5563;
            border-radius: 10px;
            padding: 10px 4px;
            font-size: 0.85rem;
            font-weight: 800;
            line-height: 1.1rem;
        }

        .month-nav-row a:hover {
            background: #1f2937;
            border-color: #6b7280;
        }

        @media (max-width: 600px) {
            .month-nav-row {
                gap: 4px;
            }

            .month-nav-row a {
                font-size: 0.78rem;
                padding: 9px 2px;
                border-radius: 9px;
            }
        }
        </style>

        <div class="month-nav-row">
            <a href="?nav=prev">◀</a>
            <a href="?nav=today">Oggi</a>
            <a href="?nav=next">▶</a>
        </div>
        """,
        unsafe_allow_html=True
    )
            
    render_calendario_mese(
        df,
        df_feste,
        int(st.session_state.cal_anno),
        int(st.session_state.cal_mese)
    )


with tab2:
    st.subheader("Aggiungi festività manuale")

    giorno_festa = st.date_input("Data festività", value=date.today(), key="giorno_festa")
    descrizione_festa = st.text_input("Descrizione", placeholder="Es. Santo Patrono")

    if st.button("Aggiungi festività", width="stretch"):
        if descrizione_festa.strip() == "":
            st.warning("Inserisci una descrizione")
        else:
            nuova_festa = pd.DataFrame([{"data": pd.to_datetime(giorno_festa), "descrizione": descrizione_festa.strip()}])

            df_feste = df_feste[df_feste["data"] != pd.to_datetime(giorno_festa)]
            df_feste = pd.concat([df_feste, nuova_festa], ignore_index=True)
            df_feste = df_feste.sort_values("data")

            salva_feste_manuali(df_feste)

            riepilogo, df_valido = calcola_riepilogo(df, df_feste)
            salva_riepilogo(riepilogo)

            st.success("Festività aggiunta")
            st.rerun()

    st.divider()

    st.subheader("Festività manuali")

    if df_feste.empty:
        st.info("Nessuna festività manuale inserita")
    else:
        st.dataframe(df_feste, width="stretch", hide_index=True)

        opzioni_feste = (df_feste["data"].dt.strftime("%Y-%m-%d") + " - " + df_feste["descrizione"]).tolist()
        feste_da_eliminare = st.multiselect("Festività da eliminare", options=opzioni_feste)

        if st.button("Elimina selezionate", width="stretch"):
            date_da_eliminare = [x.split(" - ")[0] for x in feste_da_eliminare]
            df_feste = df_feste[~df_feste["data"].dt.strftime("%Y-%m-%d").isin(date_da_eliminare)]

            salva_feste_manuali(df_feste)

            riepilogo, df_valido = calcola_riepilogo(df, df_feste)
            salva_riepilogo(riepilogo)

            st.success("Festività eliminate")
            st.rerun()


with tab3:
    st.subheader("Riepilogo trimestrale")

    if riepilogo.empty:
        st.info("Nessun giorno lavorato valido")
    else:
        st.dataframe(riepilogo, width="stretch", hide_index=True)

    st.divider()

    if not df.empty:
        anni_disponibili = sorted(df["data"].dt.year.dropna().unique(), reverse=True)
    else:
        anni_disponibili = [date.today().year]

    anno_selezionato = st.selectbox("Anno da esportare", anni_disponibili)
    df_anno = df[df["data"].dt.year == anno_selezionato].copy()

    csv_export = df_anno.to_csv(index=False).encode("utf-8")

    excel_formattato = genera_excel_formattato(df_anno, df_feste, riepilogo, anno=int(anno_selezionato))
    pdf_presenze = genera_pdf_presenze(df_anno, df_feste, anno=int(anno_selezionato))

    st.download_button(
        "Scarica CSV",
        data=csv_export,
        file_name=f"smart_calendar_{anno_selezionato}.csv",
        mime="text/csv",
        width="stretch"
    )

    st.download_button(
        "Scarica Excel",
        data=excel_formattato,
        file_name=f"smart_calendar_{anno_selezionato}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        width="stretch"
    )

    st.download_button(
        "Scarica PDF",
        data=pdf_presenze,
        file_name=f"smart_calendar_{anno_selezionato}_presenze.pdf",
        mime="application/pdf",
        width="stretch"
    )
